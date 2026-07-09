"""
Exports the "listino da usare" sheet from the pricing Excel as its own
standalone file. This is the file stages 4 (ddt) and 5 (ozanam) read as
their price list.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from wecosoft.config import ExtractListinoConfig

SHEET_NAME = "listino da usare"


def run(cfg: ExtractListinoConfig) -> dict:
    xls = pd.ExcelFile(cfg.input_xlsx)

    if SHEET_NAME not in xls.sheet_names:
        raise ValueError(f"Non trovo il foglio '{SHEET_NAME}' nel file: {cfg.input_xlsx}")

    listino = pd.read_excel(cfg.input_xlsx, sheet_name=SHEET_NAME)

    with pd.ExcelWriter(cfg.output_xlsx, engine="openpyxl") as writer:
        listino.to_excel(writer, index=False, sheet_name=SHEET_NAME)

        wb = writer.book
        ws = wb[SHEET_NAME]

        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="D9EAD3")

        widths = {"A": 42, "B": 14, "C": 34, "D": 34}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(row=1, column=col_idx).value).lower()

            if "prezzo" in header:
                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    print("File generato:", cfg.output_xlsx)
    print("Colonne:", listino.columns.tolist())

    return {"output_xlsx": cfg.output_xlsx, "n_righe": int(len(listino))}
