"""
Applies your product taxonomy CSV to the scraper's price-history Excel and
computes the sellable price list.

Input:
- The Excel produced by wecosoft.scraper (must have a PREV_matrix sheet)
- A taxonomy CSV with columns for macro category (Frutta / Verdura /
  Prodotti secchi), raggruppamento, prodotto, qualita, certificazione and
  alias. Column headers are matched loosely (case/accent-insensitive,
  ignoring any "[type]" or instructional text after the first line or
  bracket) so exports from tools like Airtable work as-is.

Output sheets:
- Tassonomia_Eco_pricing
- LISTINO COMPLETO
- listino da usare   (referenza, prezzo, prezzo -X%, prezzo +Y%)
- frutta / verdura / prodotti secchi
- Analisi andamento prezzi
- Audit_tassonomia / Audit_summary
"""

from __future__ import annotations

import json
import re
import unicodedata

import numpy as np
import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wecosoft.config import PricingConfig

MACRO_ORDER = {
    "frutta": 1,
    "verdura": 2,
    "prodotti secchi": 3,
}

MACRO_SHEETS = list(MACRO_ORDER.keys())

REQUIRED_TAX_COLS = [
    "macro_categoria",
    "raggruppamento",
    "prodotto",
    "qualita",
    "certificazione",
    "alias",
]

# Maps the loose/verbose headers used by taxonomy export tools (e.g.
# Airtable's "Macro-categoria * [select]\nFrutta / Verdura / ...") onto the
# canonical column names this module works with.
HEADER_ALIASES = {
    "macro_categoria": ["macro-categoria", "macro categoria"],
    "raggruppamento": ["raggruppamento"],
    "prodotto": ["prodotto"],
    "qualita": ["qualita"],
    "certificazione": ["certificazioni", "certificazione"],
    "alias": ["alias"],
}

QUALITY_NOISE = {
    "",
    "n c",
    "nc",
    "n c.",
    "n.c",
    "n.c.",
    "i",
    "ii",
    "extra",
    "non classificato",
    "non classificata",
}

OLD_GENERATED_SHEETS_TO_REPLACE = {
    "Tassonomia_Eco_pricing",
    "LISTINO COMPLETO",
    "listino da usare",
    "frutta",
    "frutta esotica",
    "verdura",
    "ortaggi e verdure",
    "prodotti secchi",
    "Analisi andamento prezzi",
    "Audit_tassonomia",
    "Audit_summary",
}


# ------------------------------------------------------------
# Normalization helpers
# ------------------------------------------------------------

def strip_accents(s) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))


def norm(s) -> str:
    if pd.isna(s):
        return ""

    s = strip_accents(str(s)).lower()
    s = s.replace("’", "'")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()

    return s


def parse_json_array(x):
    if pd.isna(x) or str(x).strip() == "":
        return []

    if isinstance(x, list):
        return x

    try:
        parsed = json.loads(str(x))
        return parsed if isinstance(parsed, list) else []
    except Exception:
        return []


def json_array_clean(x):
    values = parse_json_array(x)
    values = [str(v).strip() for v in values if str(v).strip()]
    return json.dumps(values, ensure_ascii=False)


def pluralize_word(w):
    w = norm(w)

    if len(w) <= 2:
        return w

    if w.endswith("io"):
        return w[:-2] + "i"
    if w.endswith("ca"):
        return w[:-2] + "che"
    if w.endswith("ga"):
        return w[:-2] + "ghe"
    if w.endswith("cia"):
        return w[:-3] + "ce"
    if w.endswith("gia"):
        return w[:-3] + "ge"
    if w.endswith("co"):
        return w[:-2] + "chi"
    if w.endswith("go"):
        return w[:-2] + "ghi"
    if w.endswith("a"):
        return w[:-1] + "e"
    if w.endswith("o"):
        return w[:-1] + "i"
    if w.endswith("e"):
        return w[:-1] + "i"

    return w


def phrase_variants(s):
    s = norm(s)

    if not s:
        return set()

    variants = {s}
    tokens = s.split()

    if tokens:
        variants.add(" ".join(tokens[:-1] + [pluralize_word(tokens[-1])]))
        variants.add(" ".join([pluralize_word(tokens[0])] + tokens[1:]))
        variants.add(" ".join([pluralize_word(t) for t in tokens]))

    return {v for v in variants if v}


def all_tokens_present(text_norm, phrase_norm):
    tokens = [t for t in norm(phrase_norm).split() if len(t) > 1]
    words = set(text_norm.split())

    return bool(tokens) and all(t in words for t in tokens)


def detect_certs(ref):
    nr = norm(ref)
    out = []

    if re.search(r"\b(i g p|igp)\b", nr):
        out.append("igp")
    if re.search(r"\b(d o p|dop)\b", nr):
        out.append("dop")
    if re.search(r"\b(bio|biologico|biologica)\b", nr):
        out.append("bio")
    if re.search(r"\b(stg)\b", nr):
        out.append("stg")

    return out


def split_caat_ref(ref):
    return [p.strip() for p in str(ref).split(" - ")]


def make_nome_referenza_eco(row):
    prodotto = str(row["prodotto"]).strip()
    qualita = str(row["qualita"]).strip()

    if qualita.lower() in ["nan", "none"]:
        qualita = ""

    certs = parse_json_array(row["certificazione"])

    parts = [prodotto]

    if qualita:
        parts.append(qualita)

    if certs:
        parts.append(" ".join([c.upper() for c in certs]))

    return " ".join(parts).strip()


def segment_matches(segment_norm, variants):
    if not segment_norm or not variants:
        return False

    if segment_norm in variants:
        return True

    for v in variants:
        if len(v) >= 4:
            if segment_norm == v:
                return True
            if segment_norm.startswith(v + " "):
                return True
            if v.startswith(segment_norm + " "):
                return True

    return False


def _normalize_header_cell(col) -> str:
    key = str(col).split("\n")[0].split("[")[0]
    key = key.replace("*", "")
    key = strip_accents(key).strip().lower()
    return key


def map_taxonomy_headers(columns) -> dict:
    """
    Matches loosely-formatted source headers (extra instructional text,
    accents, "[type]" suffixes) onto the canonical column names.
    """
    colmap = {}

    for col in columns:
        key = _normalize_header_cell(col)

        for target, aliases in HEADER_ALIASES.items():
            if target in colmap.values():
                continue
            if any(key.startswith(strip_accents(a)) for a in aliases):
                colmap[col] = target
                break

    return colmap


def load_taxonomy(taxonomy_csv: str) -> pd.DataFrame:
    taxonomy = pd.read_csv(taxonomy_csv, encoding="utf-8")
    taxonomy = taxonomy.rename(columns=map_taxonomy_headers(taxonomy.columns))

    missing_tax_cols = [c for c in REQUIRED_TAX_COLS if c not in taxonomy.columns]
    if missing_tax_cols:
        raise ValueError(f"Colonne mancanti nella tassonomia CSV: {missing_tax_cols}")

    taxonomy = taxonomy.copy()

    for c in REQUIRED_TAX_COLS:
        taxonomy[c] = taxonomy[c].fillna("").astype(str).str.strip()

    taxonomy = taxonomy[taxonomy["prodotto"] != ""].reset_index(drop=True)
    taxonomy["macro_categoria"] = taxonomy["macro_categoria"].str.lower()

    allowed_macro = set(MACRO_ORDER.keys())
    invalid_macro = sorted(set(taxonomy["macro_categoria"].dropna()) - allowed_macro)

    if invalid_macro:
        raise ValueError("La tassonomia contiene macro-categorie non ammesse: " + ", ".join(invalid_macro))

    taxonomy["certificazione"] = taxonomy["certificazione"].apply(json_array_clean)
    taxonomy["alias"] = taxonomy["alias"].apply(json_array_clean)
    taxonomy["TAX_ID"] = range(1, len(taxonomy) + 1)
    taxonomy["Nome referenza ECO"] = taxonomy.apply(make_nome_referenza_eco, axis=1)
    taxonomy["macro_sort"] = taxonomy["macro_categoria"].map(MACRO_ORDER)

    taxonomy["product_variants"] = taxonomy.apply(
        lambda r: phrase_variants(r["prodotto"])
        | set().union(*[phrase_variants(a) for a in parse_json_array(r["alias"])]),
        axis=1,
    )

    taxonomy["quality_variants"] = taxonomy.apply(
        lambda r: phrase_variants(r["qualita"])
        | set().union(*[phrase_variants(a) for a in parse_json_array(r["alias"])]),
        axis=1,
    )

    taxonomy["certs_list"] = taxonomy["certificazione"].apply(parse_json_array)

    return taxonomy


def match_reference(ref, taxonomy: pd.DataFrame) -> dict:
    """
    1. Try to map the CAAT reference to a specific quality.
       E.g. AVOCADO - HASS -> avocado Hass
    2. If no specific quality matches, fall back to the base record for the
       same product with empty quality.
       E.g. AVOCADO - 306-365 (12) -> avocado
    3. If there's no base record either, flag NON_MAPPATA_QUALITA.
    """
    ref = str(ref)
    ref_norm = norm(ref)

    if ref_norm.startswith("listino n"):
        return {
            "TAX_ID": np.nan,
            "match_status": "IGNORATA_FOOTER_LISTINO",
            "match_note": "Riga footer/paginazione del PDF, non è una referenza prodotto.",
            "match_score": 0,
        }

    parts = split_caat_ref(ref)

    if len(parts) == 0:
        return {
            "TAX_ID": np.nan,
            "match_status": "NON_MAPPATA",
            "match_note": "Referenza vuota o non interpretabile.",
            "match_score": 0,
        }

    product_segment = norm(parts[0])
    quality_segment = norm(parts[1]) if len(parts) > 1 else ""
    ref_certs = detect_certs(ref)

    product_candidates = taxonomy[
        taxonomy["product_variants"].apply(lambda variants: segment_matches(product_segment, variants))
    ].copy()

    if product_candidates.empty:
        return {
            "TAX_ID": np.nan,
            "match_status": "NON_MAPPATA",
            "match_note": f"Nessun prodotto della tassonomia corrisponde al segmento CAAT: '{parts[0]}'.",
            "match_score": 0,
        }

    specific_candidates = []
    base_candidates = []

    for _, row in product_candidates.iterrows():
        qualita = str(row["qualita"]).strip()
        row_certs = row["certs_list"]

        if row_certs:
            if not all(c in ref_certs for c in row_certs):
                continue

        cert_score = 50 if row_certs else 0

        if qualita:
            matched_quality = False

            if quality_segment not in QUALITY_NOISE:
                if segment_matches(quality_segment, row["quality_variants"]):
                    matched_quality = True

            if not matched_quality:
                for v in row["quality_variants"]:
                    if v and all_tokens_present(ref_norm, v):
                        matched_quality = True
                        break

            if matched_quality:
                score = 100 + cert_score + len(norm(qualita)) + len(norm(row["prodotto"])) / 10
                specific_candidates.append((score, row))
        else:
            score = 10 + cert_score + len(norm(row["prodotto"])) / 10
            base_candidates.append((score, row))

    if specific_candidates:
        candidates = sorted(specific_candidates, key=lambda x: x[0], reverse=True)
        match_status = "OK"
        match_note = ""
    elif base_candidates:
        candidates = sorted(base_candidates, key=lambda x: x[0], reverse=True)
        match_status = "OK"
        match_note = "Mappata su record base senza qualità perché nessuna qualità specifica è stata riconosciuta."
    else:
        return {
            "TAX_ID": np.nan,
            "match_status": "NON_MAPPATA_QUALITA",
            "match_note": (
                f"Prodotto CAAT '{parts[0]}' trovato nella tassonomia, "
                f"ma qualità/tipologia '{parts[1] if len(parts) > 1 else ''}' non mappata "
                f"e nessun record base senza qualità disponibile."
            ),
            "match_score": 0,
        }

    top_score = candidates[0][0]
    top_rows = [r for score, r in candidates if abs(score - top_score) < 0.00001]

    top_keys = {
        (r["macro_categoria"], r["raggruppamento"], r["prodotto"], r["qualita"], r["certificazione"])
        for r in top_rows
    }

    if len(top_keys) > 1:
        return {
            "TAX_ID": np.nan,
            "match_status": "AMBIGUA",
            "match_note": "Più record tassonomici candidati: "
            + " | ".join([make_nome_referenza_eco(r) for r in top_rows[:8]]),
            "match_score": top_score,
        }

    best = top_rows[0]

    return {
        "TAX_ID": best["TAX_ID"],
        "match_status": match_status,
        "match_note": match_note,
        "match_score": top_score,
    }


def run(cfg: PricingConfig) -> dict:
    xls = pd.ExcelFile(cfg.input_xlsx)

    if "PREV_matrix" not in xls.sheet_names:
        raise ValueError("Non trovo il foglio 'PREV_matrix' nel file Excel.")

    sheets = {name: pd.read_excel(cfg.input_xlsx, sheet_name=name) for name in xls.sheet_names}
    matrix = sheets["PREV_matrix"].copy()

    if "REFERENZA_COMPLETA" not in matrix.columns:
        raise ValueError("Nel foglio PREV_matrix manca la colonna 'REFERENZA_COMPLETA'.")

    taxonomy = load_taxonomy(cfg.taxonomy_csv)

    match_rows = []
    for ref in matrix["REFERENZA_COMPLETA"].astype(str):
        m = match_reference(ref, taxonomy)
        m["REFERENZA_COMPLETA"] = ref
        match_rows.append(m)

    match_df = pd.DataFrame(match_rows)

    date_cols = sorted([c for c in matrix.columns if c != "REFERENZA_COMPLETA"])
    for c in date_cols:
        matrix[c] = pd.to_numeric(matrix[c], errors="coerce")

    tax_info_cols = [
        "TAX_ID",
        "macro_categoria",
        "raggruppamento",
        "prodotto",
        "qualita",
        "certificazione",
        "alias",
        "Nome referenza ECO",
    ]
    tax_info = taxonomy[tax_info_cols].copy()

    matched_matrix = matrix.merge(match_df, on="REFERENZA_COMPLETA", how="left").merge(
        tax_info, on="TAX_ID", how="left"
    )

    latest_dates = date_cols[-cfg.last_n_listini :]

    latest_long = matched_matrix[matched_matrix["match_status"] == "OK"].melt(
        id_vars=["TAX_ID"], value_vars=latest_dates, var_name="DATA_LISTINO", value_name="PREV"
    )
    latest_long["PREV"] = pd.to_numeric(latest_long["PREV"], errors="coerce")

    # Direct prices from mapped CAAT references
    price_summary_direct = (
        latest_long.dropna(subset=["PREV"])
        .groupby("TAX_ID", as_index=False)
        .agg(
            PREZZO_BASE_MIN_ULTIMA_SETTIMANA=("PREV", "min"),
            N_PREZZI_USATI_ULTIMA_SETTIMANA=("PREV", "count"),
            DATE_LISTINI_USATI=("DATA_LISTINO", lambda x: ", ".join(sorted(set(map(str, x))))),
        )
    )
    price_summary_direct["METODO_PREZZO"] = "diretto_da_referenze_caat"

    # Prices for quality-less base records: min of direct + quality variants
    tax_price_tmp = taxonomy[
        ["TAX_ID", "Nome referenza ECO", "macro_categoria", "raggruppamento", "prodotto", "qualita", "certificazione", "alias"]
    ].copy()

    tax_price_tmp = tax_price_tmp.merge(price_summary_direct, on="TAX_ID", how="left")
    tax_price_tmp["qualita_norm"] = tax_price_tmp["qualita"].fillna("").astype(str).str.strip()

    family_price_summary = (
        tax_price_tmp.dropna(subset=["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"])
        .groupby(["macro_categoria", "raggruppamento", "prodotto"], as_index=False)
        .agg(
            PREZZO_BASE_MIN_FAMIGLIA=("PREZZO_BASE_MIN_ULTIMA_SETTIMANA", "min"),
            N_RECORD_CON_PREZZO_FAMIGLIA=("TAX_ID", "count"),
            REFERENZE_FAMIGLIA_USATE=("Nome referenza ECO", lambda x: ", ".join(sorted(set(map(str, x))))),
        )
    )

    base_rows = tax_price_tmp[tax_price_tmp["qualita_norm"] == ""].copy()
    base_rows_family_price = base_rows.merge(
        family_price_summary, on=["macro_categoria", "raggruppamento", "prodotto"], how="left"
    )
    base_rows_family_price = base_rows_family_price[
        base_rows_family_price["PREZZO_BASE_MIN_FAMIGLIA"].notna()
    ].copy()

    price_summary_base = base_rows_family_price[
        ["TAX_ID", "PREZZO_BASE_MIN_FAMIGLIA", "N_RECORD_CON_PREZZO_FAMIGLIA", "REFERENZE_FAMIGLIA_USATE"]
    ].rename(
        columns={
            "PREZZO_BASE_MIN_FAMIGLIA": "PREZZO_BASE_MIN_ULTIMA_SETTIMANA",
            "N_RECORD_CON_PREZZO_FAMIGLIA": "N_PREZZI_USATI_ULTIMA_SETTIMANA",
            "REFERENZE_FAMIGLIA_USATE": "DATE_LISTINI_USATI",
        }
    )
    price_summary_base["METODO_PREZZO"] = "record_base_da_minimo_famiglia_incluso_diretto_e_varianti"

    base_tax_ids = set(
        taxonomy[taxonomy["qualita"].fillna("").astype(str).str.strip() == ""]["TAX_ID"].astype(int).tolist()
    )
    price_summary_direct_non_base = price_summary_direct[~price_summary_direct["TAX_ID"].isin(base_tax_ids)].copy()

    price_summary = pd.concat([price_summary_direct_non_base, price_summary_base], ignore_index=True).reset_index(
        drop=True
    )

    price_summary["PREZZO_RISTORANTE_PLUS_15"] = (
        price_summary["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"] * (1 + cfg.margine_ristorante)
    ).round(2)

    taxonomy_counts = (
        matched_matrix[matched_matrix["match_status"] == "OK"]
        .groupby("TAX_ID", as_index=False)
        .agg(N_REFERENZE_CAAT_MAPPATE=("REFERENZA_COMPLETA", "nunique"))
    )

    taxonomy_pricing = matched_matrix.merge(price_summary, on="TAX_ID", how="left").merge(
        taxonomy_counts, on="TAX_ID", how="left"
    )
    taxonomy_pricing["N_REFERENZE_CAAT_MAPPATE"] = taxonomy_pricing["N_REFERENZE_CAAT_MAPPATE"].fillna(0).astype(int)

    taxonomy_pricing = taxonomy_pricing[
        [
            "REFERENZA_COMPLETA",
            "TAX_ID",
            "Nome referenza ECO",
            "macro_categoria",
            "raggruppamento",
            "prodotto",
            "qualita",
            "certificazione",
            "alias",
            "PREZZO_BASE_MIN_ULTIMA_SETTIMANA",
            "PREZZO_RISTORANTE_PLUS_15",
            "N_REFERENZE_CAAT_MAPPATE",
            "N_PREZZI_USATI_ULTIMA_SETTIMANA",
            "DATE_LISTINI_USATI",
            "match_status",
            "match_note",
            "match_score",
        ]
        + date_cols
    ].copy()

    # LISTINO COMPLETO
    if cfg.include_taxonomy_rows_without_mercuriale_match:
        listino_completo = taxonomy[tax_info_cols + ["macro_sort"]].copy()
    else:
        mapped_tax_ids = sorted(
            matched_matrix.loc[matched_matrix["match_status"] == "OK", "TAX_ID"].dropna().astype(int).unique()
        )
        listino_completo = taxonomy[taxonomy["TAX_ID"].isin(mapped_tax_ids)][tax_info_cols + ["macro_sort"]].copy()

    listino_completo = listino_completo.merge(price_summary, on="TAX_ID", how="left").merge(
        taxonomy_counts, on="TAX_ID", how="left"
    )
    listino_completo["N_REFERENZE_CAAT_MAPPATE"] = listino_completo["N_REFERENZE_CAAT_MAPPATE"].fillna(0).astype(int)

    listino_completo = listino_completo.sort_values(
        ["macro_sort", "raggruppamento", "prodotto", "qualita"], kind="stable"
    ).reset_index(drop=True)

    listino_completo = listino_completo[
        [
            "Nome referenza ECO",
            "macro_categoria",
            "raggruppamento",
            "prodotto",
            "qualita",
            "certificazione",
            "alias",
            "PREZZO_BASE_MIN_ULTIMA_SETTIMANA",
            "PREZZO_RISTORANTE_PLUS_15",
            "N_REFERENZE_CAAT_MAPPATE",
            "N_PREZZI_USATI_ULTIMA_SETTIMANA",
            "DATE_LISTINI_USATI",
        ]
    ]

    # listino da usare (only priced rows)
    listino_da_usare = listino_completo.copy()
    listino_da_usare["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"] = pd.to_numeric(
        listino_da_usare["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"], errors="coerce"
    )
    listino_da_usare = listino_da_usare[listino_da_usare["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"].notna()].copy()

    col_sconto = f"prezzo senza selezione qualità -{cfg.sconto_senza_selezione_qualita * 100:g}%"
    col_margine = f"prezzo con selezione qualità +{cfg.margine_con_selezione_qualita * 100:g}%"

    listino_da_usare = listino_da_usare[["Nome referenza ECO", "PREZZO_BASE_MIN_ULTIMA_SETTIMANA"]].rename(
        columns={"Nome referenza ECO": "referenza", "PREZZO_BASE_MIN_ULTIMA_SETTIMANA": "prezzo"}
    )

    listino_da_usare[col_sconto] = (listino_da_usare["prezzo"] * (1 - cfg.sconto_senza_selezione_qualita)).round(2)
    listino_da_usare[col_margine] = (listino_da_usare["prezzo"] * (1 + cfg.margine_con_selezione_qualita)).round(2)
    listino_da_usare["prezzo"] = listino_da_usare["prezzo"].round(2)

    listino_da_usare = listino_da_usare.sort_values(["referenza"], kind="stable").reset_index(drop=True)

    # Per-macro-category sheets
    macro_outputs = {}
    for macro in MACRO_SHEETS:
        subset = listino_completo[
            listino_completo["macro_categoria"].astype(str).str.strip().str.lower() == macro
        ].copy()

        media_row = {col: "" for col in subset.columns}
        media_row["Nome referenza ECO"] = "media"
        media_row["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"] = pd.to_numeric(
            subset["PREZZO_BASE_MIN_ULTIMA_SETTIMANA"], errors="coerce"
        ).mean()
        media_row["PREZZO_RISTORANTE_PLUS_15"] = pd.to_numeric(
            subset["PREZZO_RISTORANTE_PLUS_15"], errors="coerce"
        ).mean()

        subset = pd.concat([subset, pd.DataFrame([media_row])], ignore_index=True)
        macro_outputs[macro] = subset

    # Price trend analysis
    long_all = matched_matrix[matched_matrix["match_status"] == "OK"].melt(
        id_vars=[
            "TAX_ID",
            "Nome referenza ECO",
            "macro_categoria",
            "raggruppamento",
            "prodotto",
            "qualita",
            "certificazione",
            "alias",
        ],
        value_vars=date_cols,
        var_name="DATA_LISTINO",
        value_name="PREV",
    )
    long_all["DATA_LISTINO"] = pd.to_datetime(long_all["DATA_LISTINO"], errors="coerce")
    long_all["PREV"] = pd.to_numeric(long_all["PREV"], errors="coerce")
    long_all = long_all.dropna(subset=["DATA_LISTINO", "PREV"]).copy()

    daily_min = long_all.groupby(
        [
            "TAX_ID",
            "Nome referenza ECO",
            "macro_categoria",
            "raggruppamento",
            "prodotto",
            "qualita",
            "certificazione",
            "alias",
            "DATA_LISTINO",
        ],
        as_index=False,
    ).agg(PREZZO_PREV_MIN_GIORNO=("PREV", "min"))

    def pct_change(current, past):
        if pd.isna(current) or pd.isna(past) or past == 0:
            return np.nan
        return ((current - past) / past) * 100

    def last_value_on_or_before(series, target_date):
        s = series.dropna()
        if s.empty:
            return np.nan, pd.NaT

        eligible = s[s.index <= target_date]
        if eligible.empty:
            return np.nan, pd.NaT

        return eligible.iloc[-1], eligible.index[-1]

    analysis_rows = []
    group_cols = [
        "TAX_ID",
        "Nome referenza ECO",
        "macro_categoria",
        "raggruppamento",
        "prodotto",
        "qualita",
        "certificazione",
        "alias",
    ]

    detail_cols = []
    for col in ["VAR_%_PRIMO_ULTIMO", "VAR_%_1_MESE", "VAR_%_6_MESI", "VAR_%_1_ANNO", "VAR_%_2_ANNI", "VAR_%_3_ANNI"]:
        detail_cols += [f"DATA_USATA_{col}", f"PREZZO_USATO_{col}"]

    main_cols = group_cols[1:] + [
        "VAR_%_PRIMO_ULTIMO",
        "VAR_%_1_MESE",
        "VAR_%_6_MESI",
        "VAR_%_1_ANNO",
        "VAR_%_2_ANNI",
        "VAR_%_3_ANNI",
        "DATA_ULTIMO_PREZZO_USATO",
        "PREZZO_ULTIMO_USATO",
    ]

    if not daily_min.empty:
        last_date = daily_min["DATA_LISTINO"].max()
        first_date = daily_min["DATA_LISTINO"].min()

        targets = {
            "VAR_%_PRIMO_ULTIMO": first_date,
            "VAR_%_1_MESE": last_date - relativedelta(months=1),
            "VAR_%_6_MESI": last_date - relativedelta(months=6),
            "VAR_%_1_ANNO": last_date - relativedelta(years=1),
            "VAR_%_2_ANNI": last_date - relativedelta(years=2),
            "VAR_%_3_ANNI": last_date - relativedelta(years=3),
        }

        for keys, group in daily_min.groupby(group_cols):
            row_base = dict(zip(group_cols, keys))
            ts = group.sort_values("DATA_LISTINO").set_index("DATA_LISTINO")["PREZZO_PREV_MIN_GIORNO"]

            current_price, current_date_used = last_value_on_or_before(ts, last_date)

            row = {
                **row_base,
                "DATA_ULTIMO_PREZZO_USATO": current_date_used,
                "PREZZO_ULTIMO_USATO": current_price,
            }

            for col_name, target_date in targets.items():
                past_price, past_date_used = last_value_on_or_before(ts, target_date)
                row[col_name] = pct_change(current_price, past_price)
                row[f"DATA_USATA_{col_name}"] = past_date_used
                row[f"PREZZO_USATO_{col_name}"] = past_price

            analysis_rows.append(row)

    analisi = pd.DataFrame(analysis_rows)

    if analisi.empty:
        analisi = pd.DataFrame(columns=main_cols)
    else:
        analisi = analisi[main_cols + detail_cols].copy()
        analisi["macro_sort"] = analisi["macro_categoria"].map(MACRO_ORDER)
        analisi = (
            analisi.sort_values(["macro_sort", "raggruppamento", "prodotto", "qualita"], kind="stable")
            .drop(columns=["macro_sort"])
            .reset_index(drop=True)
        )

    # Taxonomy match audit
    audit_tassonomia = match_df.merge(tax_info, on="TAX_ID", how="left")
    audit_tassonomia = audit_tassonomia[
        [
            "REFERENZA_COMPLETA",
            "match_status",
            "match_note",
            "match_score",
            "TAX_ID",
            "Nome referenza ECO",
            "macro_categoria",
            "raggruppamento",
            "prodotto",
            "qualita",
            "certificazione",
            "alias",
        ]
    ].copy()

    audit_summary = (
        audit_tassonomia.groupby("match_status", as_index=False)
        .agg(n_referenze=("REFERENZA_COMPLETA", "count"))
        .sort_values("n_referenze", ascending=False)
    )

    _write_excel(
        cfg,
        sheets,
        taxonomy_pricing,
        listino_completo,
        listino_da_usare,
        macro_outputs,
        analisi,
        audit_tassonomia,
        audit_summary,
    )

    print("File generato:", cfg.output_xlsx)
    print("\nRiepilogo match tassonomia:")
    print(audit_summary.to_string(index=False))

    non_mapped = audit_tassonomia[
        audit_tassonomia["match_status"].isin(["NON_MAPPATA", "NON_MAPPATA_QUALITA", "AMBIGUA"])
    ]
    if not non_mapped.empty:
        print(f"\n{len(non_mapped)} referenze non mappate o ambigue: vedi foglio 'Audit_tassonomia'.")

    return {
        "output_xlsx": cfg.output_xlsx,
        "n_referenze_prezzate": int(len(listino_da_usare)),
        "n_non_mappate_o_ambigue": int(len(non_mapped)),
    }


def _write_excel(
    cfg: PricingConfig,
    sheets,
    taxonomy_pricing,
    listino_completo,
    listino_da_usare,
    macro_outputs,
    analisi,
    audit_tassonomia,
    audit_summary,
):
    with pd.ExcelWriter(cfg.output_xlsx, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            if sheet_name in OLD_GENERATED_SHEETS_TO_REPLACE:
                continue
            df.to_excel(writer, index=False, sheet_name=sheet_name[:31])

        taxonomy_pricing.to_excel(writer, index=False, sheet_name="Tassonomia_Eco_pricing")
        listino_completo.to_excel(writer, index=False, sheet_name="LISTINO COMPLETO")
        listino_da_usare.to_excel(writer, index=False, sheet_name="listino da usare")

        for sheet_name, df in macro_outputs.items():
            df.to_excel(writer, index=False, sheet_name=sheet_name[:31])

        analisi.to_excel(writer, index=False, sheet_name="Analisi andamento prezzi")
        audit_tassonomia.to_excel(writer, index=False, sheet_name="Audit_tassonomia")
        audit_summary.to_excel(writer, index=False, sheet_name="Audit_summary")

        wb = writer.book

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"

            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor="D9EAD3")

            for col_idx in range(1, min(ws.max_column, 20) + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = 24

            headers = {ws.cell(row=1, column=col_idx).value: col_idx for col_idx in range(1, ws.max_column + 1)}

            for header, col_idx in headers.items():
                letter = get_column_letter(col_idx)

                if header in ["REFERENZA_COMPLETA"]:
                    ws.column_dimensions[letter].width = 110
                elif header in ["Nome referenza ECO", "referenza"]:
                    ws.column_dimensions[letter].width = 42
                elif header in ["match_note"]:
                    ws.column_dimensions[letter].width = 70
                elif header in ["alias", "certificazione"]:
                    ws.column_dimensions[letter].width = 26
                elif header and "PREZZO" in str(header).upper():
                    ws.column_dimensions[letter].width = 22
                    for row_idx in range(2, ws.max_row + 1):
                        ws.cell(row=row_idx, column=col_idx).number_format = "0.00"
                elif header and str(header).startswith("VAR_%"):
                    ws.column_dimensions[letter].width = 18
                    for row_idx in range(2, ws.max_row + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        if isinstance(cell.value, (int, float)):
                            cell.value = cell.value / 100
                        cell.number_format = "0.00%"

            if ws.title == "listino da usare":
                widths = {"A": 42, "B": 14, "C": 32, "D": 32}
                for col, width in widths.items():
                    ws.column_dimensions[col].width = width

                for row_idx in range(2, ws.max_row + 1):
                    for col_idx in [2, 3, 4]:
                        ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

            if ws.title in MACRO_SHEETS:
                for row in ws.iter_rows(min_row=2):
                    if str(row[0].value).strip().lower() == "media":
                        for cell in row:
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="FFF2CC")
